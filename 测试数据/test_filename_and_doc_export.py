# -*- coding: utf-8 -*-
import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "exam_converter", ROOT / "E学院试题格式互转.py")
converter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(converter)


class FilenameAndDocExportTests(unittest.TestCase):
    def test_output_suffix_matches_target_format(self):
        old_text = "1、旧版题目\nA.选项\nB.选项\n【答案】A"
        new_text = "【单选题开始】\n1、新版题目\nA.选项\nB.选项\n正确答案：A\n【单选题结束】"

        self.assertEqual(
            converter.decide_direction(old_text, converter.DIR_AUTO),
            (converter.FMT_OLD, converter.FMT_NEW, "_2026"),
        )
        self.assertEqual(
            converter.decide_direction(new_text, converter.DIR_AUTO),
            (converter.FMT_NEW, converter.FMT_OLD, "_2023"),
        )
        self.assertEqual(
            converter.decide_direction(old_text, converter.DIR_TO_OLD),
            (converter.FMT_OLD, converter.FMT_OLD, "_2023"),
        )
        self.assertEqual(
            converter.decide_direction(new_text, converter.DIR_TO_NEW),
            (converter.FMT_NEW, converter.FMT_NEW, "_2026"),
        )

    def test_output_filename_replaces_source_version_suffix(self):
        with tempfile.TemporaryDirectory() as output_dir:
            source = str(Path(output_dir) / "sample_2023.doc")
            path, renamed = converter.output_path_for(source, "_2026", output_dir)

        self.assertEqual(Path(path).name, "sample_2026.doc")
        self.assertFalse(renamed)

    def test_doc_export_fails_without_creating_docx_fallback(self):
        class FailingBridge:
            @staticmethod
            def docx_to_doc(_docx_path, _doc_path):
                return False

        with tempfile.TemporaryDirectory() as output_dir:
            doc_path = Path(output_dir) / "sample_2026.doc"
            with self.assertRaisesRegex(RuntimeError, r"无法导出 \.doc"):
                converter.write_document(["题目内容"], str(doc_path), FailingBridge())

            self.assertFalse(doc_path.exists())
            self.assertFalse(doc_path.with_suffix(".docx").exists())

    def test_docx_import_from_paragraphs_and_tables_and_docx_export(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "sample_2023.docx"
            document = Document()
            document.add_paragraph("1、DOCX导入题")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "A.选项甲"
            table.cell(0, 1).text = "B.选项乙"
            document.add_paragraph("正确答案：A")
            document.save(source)

            text, engine = converter.read_document(str(source))
            self.assertEqual(engine, "python-docx")
            self.assertIn("DOCX导入题", text)
            self.assertIn("A.选项甲", text)
            self.assertIn("B.选项乙", text)

            output, stats, _, _, write_engine = converter.convert_file(
                str(source),
                converter.DIR_AUTO,
                temp,
                output_format=".docx",
            )

            self.assertEqual(Path(output).name, "sample_2026.docx")
            self.assertEqual(write_engine, "python-docx")
            self.assertEqual(stats["single"], 1)
            exported_text, _ = converter.read_document(output)
            self.assertIn("【单选题开始】", exported_text)
            self.assertIn("DOCX导入题", exported_text)

    def test_invalid_docx_reports_read_error_without_binary_fallback(self):
        with tempfile.TemporaryDirectory() as temp:
            invalid_docx = Path(temp) / "invalid.docx"
            invalid_docx.write_bytes(b"not a valid docx")

            with self.assertRaisesRegex(RuntimeError, "无法读取 DOCX"):
                converter.read_document(str(invalid_docx))


if __name__ == "__main__":
    unittest.main()
